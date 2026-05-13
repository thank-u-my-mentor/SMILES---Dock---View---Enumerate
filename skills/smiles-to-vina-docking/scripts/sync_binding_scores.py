#!/usr/bin/env python
"""Merge local official binding scores into dock_history.csv and analyze trends."""

from __future__ import annotations

import argparse
import csv
import math
import re
from pathlib import Path

from openpyxl import load_workbook

import dock_utils as dl


ANALYSIS_FIELDS = [
    "seq_id",
    "nickname",
    "input_smiles",
    "canonical_smiles",
    "official_binding_score",
    "affinity_kcal_mol",
    "mode_count",
    "inner_rmsd",
    "inner_cluster_fraction",
    "whole_rmsd",
    "hbond_count",
    "hydrophobic_count",
    "vdw_contact_count",
    "pi_contact_count",
    "ch_pi_count",
    "contact_residue_count_4a",
    "contact_residue_count_5a",
    "ligand_atom_contact_fraction_4a",
    "min_ligand_receptor_distance",
    "closest_residues",
    "pose_path",
    "log_path",
]

CORRELATION_FIELDS = ["metric", "n", "pearson_r", "spearman_r"]
MISSING_FIELDS = ["sheet_index", "row", "mol_smiles", "canonical_smiles", "official_binding_score"]


def norm_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def parse_float(value: object) -> float | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def read_score_rows(xlsx_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    rows: list[dict[str, str]] = []
    smiles_aliases = {"molsmiles", "smiles", "moleculesmiles", "ligandsmiles"}
    score_aliases = {"bindingscores", "bindingscore", "officialbindingscore", "score"}
    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        values = sheet.iter_rows(values_only=True)
        try:
            headers = next(values)
        except StopIteration:
            continue
        header_map = {norm_header(value): index for index, value in enumerate(headers)}
        smiles_idx = next((header_map[name] for name in smiles_aliases if name in header_map), None)
        score_idx = next((header_map[name] for name in score_aliases if name in header_map), None)
        if smiles_idx is None or score_idx is None:
            continue
        for row_number, row in enumerate(values, start=2):
            smiles = dl.clean_smiles(row[smiles_idx] if smiles_idx < len(row) else "")
            score = parse_float(row[score_idx] if score_idx < len(row) else None)
            if not smiles or score is None:
                continue
            rows.append(
                {
                    "sheet": sheet.title,
                    "sheet_index": str(sheet_index),
                    "row": str(row_number),
                    "mol_smiles": smiles,
                    "canonical_smiles": dl.canonical(smiles) or "",
                    "official_binding_score": str(score),
                }
            )
    return rows


def write_csv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def best_mode(row: dict[str, str]) -> dl.PoseMode | None:
    pose_path = Path(row.get("pose_path", "")).expanduser()
    modes = dl.parse_pdbqt_modes(pose_path)
    if not modes:
        return None
    return min(modes, key=lambda mode: mode.affinity if mode.affinity is not None else float("inf"))


def residue_label(atom: dict[str, str | float]) -> str:
    return f"{atom.get('residue_name', '')}{atom.get('residue_id', '')}"


def pocket_metrics(row: dict[str, str], receptor_atoms: list[dict[str, str | float]]) -> dict[str, object]:
    mode = best_mode(row)
    if mode is None or not receptor_atoms:
        return {}
    min_distance = float("inf")
    contacts4: set[str] = set()
    contacts5: set[str] = set()
    contacted_ligand_atoms: set[int] = set()
    residue_distances: dict[str, float] = {}
    for i, ligand_atom in enumerate(mode.atoms):
        atom_min = float("inf")
        for receptor_atom in receptor_atoms:
            distance = math.sqrt(dl.squared_distance(ligand_atom, receptor_atom))
            label = residue_label(receptor_atom)
            if distance < min_distance:
                min_distance = distance
            if distance < atom_min:
                atom_min = distance
            if label:
                previous = residue_distances.get(label)
                if previous is None or distance < previous:
                    residue_distances[label] = distance
            if distance <= 4.0 and label:
                contacts4.add(label)
                contacted_ligand_atoms.add(i)
            if distance <= 5.0 and label:
                contacts5.add(label)
        if atom_min <= 4.0:
            contacted_ligand_atoms.add(i)
    closest = sorted(residue_distances.items(), key=lambda item: item[1])[:8]
    return {
        "contact_residue_count_4a": len(contacts4),
        "contact_residue_count_5a": len(contacts5),
        "ligand_atom_contact_fraction_4a": round(len(contacted_ligand_atoms) / len(mode.atoms), 4) if mode.atoms else "",
        "min_ligand_receptor_distance": round(min_distance, 4) if math.isfinite(min_distance) else "",
        "closest_residues": ";".join(f"{label}:{distance:.2f}" for label, distance in closest),
    }


def ranks(values: list[float]) -> list[float]:
    order = sorted(range(len(values)), key=lambda index: values[index])
    out = [0.0] * len(values)
    i = 0
    while i < len(order):
        j = i + 1
        while j < len(order) and values[order[j]] == values[order[i]]:
            j += 1
        rank = (i + j + 1) / 2.0
        for k in range(i, j):
            out[order[k]] = rank
        i = j
    return out


def pearson(x: list[float], y: list[float]) -> float | None:
    if len(x) < 3 or len(x) != len(y):
        return None
    mx = sum(x) / len(x)
    my = sum(y) / len(y)
    numerator = sum((a - mx) * (b - my) for a, b in zip(x, y))
    dx = math.sqrt(sum((a - mx) ** 2 for a in x))
    dy = math.sqrt(sum((b - my) ** 2 for b in y))
    if dx == 0.0 or dy == 0.0:
        return None
    return numerator / (dx * dy)


def correlation_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    metrics = [
        "affinity_kcal_mol",
        "inner_rmsd",
        "inner_cluster_fraction",
        "whole_rmsd",
        "hbond_count",
        "hydrophobic_count",
        "vdw_contact_count",
        "pi_contact_count",
        "ch_pi_count",
        "contact_residue_count_4a",
        "contact_residue_count_5a",
        "ligand_atom_contact_fraction_4a",
        "min_ligand_receptor_distance",
    ]
    out: list[dict[str, object]] = []
    for metric in metrics:
        x: list[float] = []
        y: list[float] = []
        for row in rows:
            score = parse_float(row.get("official_binding_score"))
            value = parse_float(row.get(metric))
            if score is None or value is None:
                continue
            x.append(value)
            y.append(score)
        pr = pearson(x, y)
        sr = pearson(ranks(x), ranks(y)) if len(x) >= 3 else None
        out.append(
            {
                "metric": metric,
                "n": len(x),
                "pearson_r": "" if pr is None else round(pr, 6),
                "spearman_r": "" if sr is None else round(sr, 6),
            }
        )
    return out


def group_scores(score_rows: list[dict[str, str]], xlsx_name: str) -> dict[str, dict[str, str]]:
    score_by_canonical: dict[str, dict[str, str]] = {}
    for score_row in score_rows:
        can = score_row.get("canonical_smiles", "")
        score = parse_float(score_row.get("official_binding_score"))
        if not can or score is None:
            continue
        existing = score_by_canonical.setdefault(
            can,
            {
                "canonical_smiles": can,
                "mol_smiles": score_row.get("mol_smiles", ""),
                "scores": "",
                "sources": "",
            },
        )
        scores = [text for text in existing["scores"].split(";") if text]
        sources = [text for text in existing["sources"].split(";") if text]
        scores.append(str(score))
        sheet_ref = score_row.get("sheet_index") or score_row.get("sheet") or ""
        sources.append(f"{xlsx_name}:sheet{sheet_ref}:row{score_row['row']}")
        existing["scores"] = ";".join(scores)
        existing["sources"] = ";".join(sources)
    return score_by_canonical


def match_scores(
    rows: list[dict[str, str]],
    score_by_canonical: dict[str, dict[str, str]],
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    history_by_canonical: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        can = row.get("canonical_smiles", "")
        if can:
            history_by_canonical.setdefault(can, []).append(row)

    matched: list[dict[str, object]] = []
    missing: list[dict[str, object]] = []
    for can, score_group in score_by_canonical.items():
        score_values = [parse_float(text) for text in score_group["scores"].split(";") if text]
        finite_scores = [score for score in score_values if score is not None]
        if not finite_scores:
            continue
        average_score = sum(finite_scores) / len(finite_scores)
        history_matches = history_by_canonical.get(can, [])
        if not history_matches:
            missing.append(
                {
                    "sheet_index": "",
                    "row": "",
                    "mol_smiles": score_group.get("mol_smiles", ""),
                    "canonical_smiles": can,
                    "official_binding_score": str(round(average_score, 8)),
                }
            )
            continue
        for row in history_matches:
            row["official_binding_score"] = str(round(average_score, 8))
            row["binding_score_source"] = score_group["sources"]
        analysis_row = min(
            history_matches,
            key=lambda row: float(row.get("affinity_kcal_mol") or 999.0),
        )
        matched.append(analysis_row)
    return matched, missing


def dock_missing_scores(args: argparse.Namespace, missing: list[dict[str, object]]) -> None:
    if not missing:
        return
    limit = args.dock_missing_limit if args.dock_missing_limit and args.dock_missing_limit > 0 else len(missing)
    to_dock = missing[:limit]
    print(f"dock_missing=true count={len(to_dock)}", flush=True)
    for index, item in enumerate(to_dock, start=1):
        smiles = str(item.get("mol_smiles") or item.get("canonical_smiles") or "")
        score = str(item.get("official_binding_score") or "")
        if not smiles:
            continue
        dock_args = argparse.Namespace(
            smiles=smiles,
            nickname=f"score_{score}" if score else "",
            ledger_dir=args.history_dir.expanduser().resolve(),
            receptor=args.receptor,
            config=args.config,
            vina=args.vina,
            meeko=args.meeko,
            obabel=args.obabel,
            max_rounds=0,
            batch_size=0,
            seed=args.seed + index,
            cpu=args.cpu,
            exhaustiveness=args.exhaustiveness,
            num_modes=args.num_modes,
            energy_range=args.energy_range,
            internal_cluster_rmsd_cutoff=args.internal_cluster_rmsd_cutoff,
            edit_mode="default",
            add_only=False,
            delete_only=False,
            shrink_only=False,
            drastic_only=False,
            deterministic_batch=True,
            redock_existing=args.redock_existing,
        )
        print(f"[dock-missing {index}/{len(to_dock)}] {smiles}", flush=True)
        dl.collect(dock_args)


def sync(args: argparse.Namespace) -> None:
    history_dir = args.history_dir.expanduser().resolve()
    history_path = dl.ledger_path(history_dir)
    output_path = args.output.expanduser().resolve() if args.output else history_dir / "binding_score_analysis.csv"
    correlations_path = args.correlations.expanduser().resolve() if args.correlations else history_dir / "binding_score_correlations.csv"
    missing_path = args.missing.expanduser().resolve() if args.missing else history_dir / "binding_score_missing_smiles.csv"
    score_rows = read_score_rows(args.xlsx.expanduser().resolve())
    score_by_canonical = group_scores(score_rows, args.xlsx.name)
    rows = dl.read_csv(history_path)
    matched, missing = match_scores(rows, score_by_canonical)

    if args.dock_missing and missing:
        dock_missing_scores(args, missing)
        rows = dl.read_csv(history_path)
        matched, missing = match_scores(rows, score_by_canonical)

    receptor_atoms: list[dict[str, str | float]] = []
    if args.receptor:
        receptor_atoms = dl.parse_receptor_atoms(args.receptor.expanduser().resolve())
    analysis_rows: list[dict[str, object]] = []
    for row in matched:
        analysis_rows.append({**row, **pocket_metrics(row, receptor_atoms)})

    if args.update_history:
        dl.write_csv(history_path, rows)
    write_csv(output_path, analysis_rows, ANALYSIS_FIELDS)
    write_csv(correlations_path, correlation_rows(analysis_rows), CORRELATION_FIELDS)
    write_csv(missing_path, missing, MISSING_FIELDS)
    print(
        f"score_rows={len(score_rows)} matched={len(matched)} missing={len(missing)} "
        f"updated_history={args.update_history}",
        flush=True,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, required=True, help="Local xlsx with mol_smiles and binding_scores columns")
    parser.add_argument("--history-dir", "--ledger-dir", dest="history_dir", type=Path, required=True)
    parser.add_argument("--receptor", type=Path, help="Optional receptor PDBQT/PDBQT-like file for ligand-pocket distance metrics")
    parser.add_argument("--config", type=Path, help="Required for --dock-missing unless saved in history_config.json")
    parser.add_argument("--vina", default="vina")
    parser.add_argument("--meeko", default="mk_prepare_ligand.py")
    parser.add_argument("--obabel", default="obabel")
    parser.add_argument("--cpu", type=int)
    parser.add_argument("--exhaustiveness", type=int)
    parser.add_argument("--num-modes", type=int, default=9)
    parser.add_argument("--energy-range", type=int)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--internal-cluster-rmsd-cutoff", type=float, default=3.0)
    parser.add_argument("--update-history", action="store_true", help="Write official_binding_score back to dock_history.csv")
    parser.add_argument("--dock-missing", action="store_true", help="Dock xlsx SMILES not yet present in dock_history.csv, then rerun score sync")
    parser.add_argument("--dock-missing-limit", type=int, default=0, help="Limit number of missing SMILES to dock; 0 means all")
    parser.add_argument("--redock-existing", action="store_true", help="Passed to dock_smiles logic when --dock-missing is used")
    parser.add_argument("--output", type=Path, help="Default: <history-dir>/binding_score_analysis.csv")
    parser.add_argument("--correlations", type=Path, help="Default: <history-dir>/binding_score_correlations.csv")
    parser.add_argument("--missing", type=Path, help="Default: <history-dir>/binding_score_missing_smiles.csv")
    args = parser.parse_args()
    sync(args)


if __name__ == "__main__":
    main()
